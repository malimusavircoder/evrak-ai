import os
import json
import base64
import shutil
import re
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import requests
from dotenv import load_dotenv

os.environ.setdefault("PYDANTIC_DISABLE_PLUGINS", "1")

from fastapi import FastAPI, Request, UploadFile, File, Form, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import create_engine, Column, Integer, String, Float, DateTime
from sqlalchemy.orm import declarative_base, sessionmaker

try:
    import pytesseract
    from PIL import Image, ImageFilter, ImageOps, UnidentifiedImageError
    from pytesseract import TesseractNotFoundError
except ImportError:  # pragma: no cover - optional runtime dependency
    pytesseract = None
    Image = None
    ImageFilter = None
    ImageOps = None
    UnidentifiedImageError = Exception
    TesseractNotFoundError = RuntimeError

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
EXPORT_DIR = BASE_DIR / "exports"
DATA_DIR = BASE_DIR / "data"
STATIC_DIR = BASE_DIR / "static"

load_dotenv(BASE_DIR / ".env", override=True, encoding="utf-8-sig")

UPLOAD_DIR.mkdir(exist_ok=True)
EXPORT_DIR.mkdir(exist_ok=True)
DATA_DIR.mkdir(exist_ok=True)
STATIC_DIR.mkdir(exist_ok=True)

CLAUDE_API_KEY = os.getenv("CLAUDE_API_KEY", "")
CLAUDE_MODEL = os.getenv("CLAUDE_MODEL", "claude-opus-4-6")
SUPABASE_URL = os.getenv("SUPABASE_URL", "").rstrip("/")
SUPABASE_ANON_KEY = os.getenv("SUPABASE_ANON_KEY", "")
SUPABASE_SECRET_KEY = os.getenv("SUPABASE_SECRET_KEY", "")
SUPABASE_BUCKET = os.getenv("SUPABASE_BUCKET", "evrak-files")
MODEL_ALIASES = {
    "claude-opus-4.6": "claude-opus-4-6",
}


def normalize_model_name(value: str) -> str:
    model_name = (value or "").strip()
    if "=" in model_name:
        model_name = model_name.split("=")[-1].strip()
    return MODEL_ALIASES.get(model_name, model_name or "claude-opus-4-6")

DATABASE_URL = os.getenv("DATABASE_URL", f"sqlite:///{DATA_DIR / 'app.db'}")
if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql+psycopg://", 1)
engine_kwargs = {}
if DATABASE_URL.startswith("sqlite"):
    engine_kwargs["connect_args"] = {"check_same_thread": False}
engine = create_engine(DATABASE_URL, **engine_kwargs)
SessionLocal = sessionmaker(bind=engine)
Base = declarative_base()

app = FastAPI(title="Evrak AI")

jinja_env = Environment(
    loader=FileSystemLoader(str(BASE_DIR / "templates")),
    autoescape=select_autoescape(["html", "xml"]),
    cache_size=0,
)

app.mount("/uploads", StaticFiles(directory=str(UPLOAD_DIR)), name="uploads")
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")


class Record(Base):
    __tablename__ = "records"
    id = Column(Integer, primary_key=True, index=True)
    musteri = Column(String, default="")
    belge_tipi = Column(String, default="")
    firma = Column(String, default="")
    vergi_no = Column(String, default="")
    tarih = Column(String, default="")
    belge_no = Column(String, default="")
    matrah = Column(Float, default=0.0)
    kdv_orani = Column(String, default="")
    kdv_tutari = Column(Float, default=0.0)
    toplam = Column(Float, default=0.0)
    aciklama = Column(String, default="")
    dosya_adi = Column(String, default="")
    dosya_yolu = Column(String, default="")
    durum = Column(String, default="Kontrol")
    created_at = Column(DateTime, default=datetime.utcnow)


class FuelRecord(Base):
    __tablename__ = "fuel_records"
    id = Column(Integer, primary_key=True, index=True)
    musteri = Column(String, default="")
    istasyon_adi = Column(String, default="")
    vergi_no = Column(String, default="")
    tarih = Column(String, default="")
    saat = Column(String, default="")
    fis_no = Column(String, default="")
    plaka = Column(String, default="")
    litre = Column(Float, default=0.0)
    birim_fiyat = Column(Float, default=0.0)
    kdv_tutari = Column(Float, default=0.0)
    toplam = Column(Float, default=0.0)
    odeme_tipi = Column(String, default="")
    aciklama = Column(String, default="")
    ocr_metin = Column(String, default="")
    dosya_adi = Column(String, default="")
    dosya_yolu = Column(String, default="")
    durum = Column(String, default="Kontrol")
    created_at = Column(DateTime, default=datetime.utcnow)


Base.metadata.create_all(bind=engine)


def safe_float(value) -> float:
    if value in (None, "", "-"):
        return 0.0
    try:
        text = str(value).replace("TL", "").replace("TRY", "").replace("%", "").strip()
        text = re.sub(r"[^\d,.\-]", "", text)
        if not text or text in {"-", ",", "."}:
            return 0.0

        last_dot = text.rfind(".")
        last_comma = text.rfind(",")

        if last_dot != -1 and last_comma != -1:
            if last_dot > last_comma:
                decimal_sep = "."
                thousands_sep = ","
            else:
                decimal_sep = ","
                thousands_sep = "."
            text = text.replace(thousands_sep, "")
            text = text.replace(decimal_sep, ".")
        elif last_comma != -1:
            digits_after = len(text) - last_comma - 1
            if digits_after == 3 and text.count(",") == 1:
                text = text.replace(",", "")
            else:
                text = text.replace(".", "")
                text = text.replace(",", ".")
        elif last_dot != -1:
            digits_after = len(text) - last_dot - 1
            if digits_after == 3 and text.count(".") == 1:
                text = text.replace(".", "")
            else:
                text = text.replace(",", "")

        return float(text)
    except Exception:
        return 0.0


def file_to_base64(filepath: Path) -> str:
    with open(filepath, "rb") as f:
        return base64.b64encode(f.read()).decode("utf-8")


def get_mime_type(filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".png"):
        return "image/png"
    if lower.endswith(".jpg") or lower.endswith(".jpeg"):
        return "image/jpeg"
    if lower.endswith(".webp"):
        return "image/webp"
    if lower.endswith(".pdf"):
        return "application/pdf"
    return "application/octet-stream"


def is_image_mime_type(mime_type: str) -> bool:
    return mime_type.startswith("image/")


def is_supabase_storage_enabled() -> bool:
    return bool(SUPABASE_URL and SUPABASE_SECRET_KEY and SUPABASE_BUCKET)


def build_public_file_url(path: str) -> str:
    if not path:
        return ""
    if path.startswith("http://") or path.startswith("https://"):
        return path
    return f"/uploads/{quote(path)}"


def sanitize_upload_filename(filename: str) -> str:
    clean_name = Path(filename).name.replace(" ", "_")
    clean_name = re.sub(r"[^A-Za-z0-9_.-]", "_", clean_name)
    return clean_name or "upload"


def upload_to_supabase_storage(filepath: Path, object_name: str, mime_type: str) -> str:
    if not is_supabase_storage_enabled():
        return build_public_file_url(object_name)

    upload_url = f"{SUPABASE_URL}/storage/v1/object/{SUPABASE_BUCKET}/{object_name}"
    headers = {
        "apikey": SUPABASE_SECRET_KEY,
        "Authorization": f"Bearer {SUPABASE_SECRET_KEY}",
        "Content-Type": mime_type,
        "x-upsert": "true",
    }
    with open(filepath, "rb") as file_handle:
        response = requests.post(upload_url, headers=headers, data=file_handle, timeout=120)
    try:
        response.raise_for_status()
    except requests.HTTPError as exc:
        detail = response.text[:500] if response.text else "Dosya depolamaya yuklenemedi."
        raise HTTPException(status_code=502, detail=detail) from exc

    return f"{SUPABASE_URL}/storage/v1/object/public/{SUPABASE_BUCKET}/{object_name}"


def should_keep_local_upload() -> bool:
    return not is_supabase_storage_enabled()


def ensure_ocr_available():
    if pytesseract is None or Image is None:
        raise HTTPException(
            status_code=500,
            detail="OCR kutuphaneleri kurulu degil. pytesseract ve Pillow gerekli.",
        )
    try:
        pytesseract.get_tesseract_version()
    except (TesseractNotFoundError, FileNotFoundError) as exc:
        raise HTTPException(
            status_code=500,
            detail="Tesseract OCR kurulu degil. Sunucuda tesseract binary gerekli.",
        ) from exc


def preprocess_fuel_image(filepath: Path):
    with Image.open(filepath) as image:
        base = ImageOps.exif_transpose(image).convert("L")
        base = ImageOps.autocontrast(base)
        base = base.filter(ImageFilter.MedianFilter(size=3))

        original = base.resize((base.width * 2, base.height * 2))
        strong = base.filter(ImageFilter.SHARPEN)
        thresholded = base.point(lambda pixel: 255 if pixel > 168 else 0)
        enlarged = strong.resize((strong.width * 2, strong.height * 2))

        return [original, strong, thresholded, enlarged]


def extract_text_with_ocr(filepath: Path) -> str:
    ensure_ocr_available()
    try:
        image_variants = preprocess_fuel_image(filepath)
    except UnidentifiedImageError as exc:
        raise HTTPException(status_code=400, detail="Yakıt fişi icin gecerli bir resim yukleyin.") from exc

    configs = [
        "--oem 3 --psm 6",
        "--oem 3 --psm 4",
        "--oem 3 --psm 11",
    ]
    best_text = ""
    best_score = -1

    for image in image_variants:
        for config in configs:
            text = pytesseract.image_to_string(image, lang="tur+eng", config=config)
            cleaned = re.sub(r"[ \t]+", " ", text)
            lines = [line.strip() for line in cleaned.splitlines() if line.strip()]
            normalized = "\n".join(lines)
            normalized_upper = normalize_ocr_text(normalized)
            score = len(lines)
            score += sum(3 for token in ("TOPLAM", "KDV", "LT X", "TARIH", "SAAT", "FIS NO", "VD:", "K. KART", "UTTS") if token in normalized_upper)
            score += 4 if re.search(r"\b\d{2}[./-]\d{2}[./-]\d{4}\b", normalized_upper) else 0
            score += 4 if re.search(r"\b\d{2}[:.]\d{2}\b", normalized_upper) else 0
            score += 4 if re.search(r"\b\d{10}\b", normalized_upper) else 0
            score += 4 if re.search(r"\b(0[1-9]|[1-7][0-9]|8[01])\s?[A-Z]{1,3}\s?\d{2,4}\b", normalized_upper) else 0
            if score > best_score:
                best_score = score
                best_text = normalized

    return best_text


def normalize_ocr_text(text: str) -> str:
    text = text.upper()
    replacements = {
        "İ": "I",
        "Ş": "S",
        "Ğ": "G",
        "Ü": "U",
        "Ö": "O",
        "Ç": "C",
        "|": "I",
        "$": "3",
        "€": "E",
        "©": "0",
        "“": "",
        "”": "",
        "'": "",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return re.sub(r"\s+", " ", text).strip()


def compact_ocr_text(text: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", normalize_ocr_text(text))


def normalize_number_token(value: str) -> str:
    token = (value or "").strip()
    token = token.replace(" ", "")
    token = token.replace("*", "")
    token = token.replace("O", "0")
    token = token.replace("S", "5") if token.count(",") == 0 and token.count(".") == 0 else token
    return token


def extract_line_amount(line: str) -> str:
    normalized_line = normalize_ocr_text(line)
    matches = re.findall(r"(\d[\d.,]{1,})", normalized_line)
    if not matches:
        return ""
    candidates = [normalize_number_token(match) for match in matches]
    candidates = [value for value in candidates if any(ch in value for ch in ",.")]
    return candidates[-1] if candidates else ""


def extract_first_match(patterns: list[str], text: str) -> str:
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            return match.group(1).strip()
    return ""


def normalize_date(value: str) -> str:
    value = (value or "").strip()
    if not value:
        return ""
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%y", "%d/%m/%y", "%d-%m-%y"):
        try:
            parsed = datetime.strptime(value, fmt)
            return parsed.strftime("%Y-%m-%d")
        except ValueError:
            continue
    return value


def normalize_time(value: str) -> str:
    value = (value or "").strip().replace(".", ":")
    match = re.search(r"(\d{2}:\d{2})", value)
    return match.group(1) if match else ""


def parse_record_date(value: str) -> datetime:
    text = (value or "").strip()
    if not text:
        return datetime.now()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%y", "%d/%m/%y", "%d-%m-%y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return datetime.now()


def build_sales_receipt_no(date_value: datetime) -> str:
    return date_value.strftime("%y%m%d") + "SF"


def safe_kdv_rate(value: str, matrah: float = 0.0, kdv_tutari: float = 0.0) -> float:
    rate = safe_float(value)
    if rate:
        return rate
    if matrah and kdv_tutari:
        return round((kdv_tutari / matrah) * 100, 2)
    return 0.0


def kdv_excluded_amount(record: Record) -> float:
    matrah = safe_float(record.matrah)
    if matrah:
        return round(matrah, 2)
    toplam = safe_float(record.toplam)
    kdv_rate = safe_kdv_rate(record.kdv_orani, 0.0, safe_float(record.kdv_tutari))
    if toplam and kdv_rate:
        return round(toplam / (1 + (kdv_rate / 100)), 2)
    return round(toplam, 2)


def pick_station_name(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for line in lines[:6]:
        upper = normalize_ocr_text(line)
        if any(token in upper for token in ("OPET", "SHELL", "BP", "PETROL OFISI", "PETROLOFISI", "AYTEMIZ", "TOTAL", "SUNPET", "LUKOIL", "ALPET", "KADOIL", "GO")):
            return line
    for line in lines[:6]:
        if "PETROL" in normalize_ocr_text(line):
            return line
    return lines[0] if lines else ""


def extract_plate(text: str) -> str:
    patterns = [
        r"\b([0-8][0-9]\s?[A-Z]{1,3}\s?\d{2,4})\b",
        r"\b([0-8][0-9][A-Z]{1,3}\d{2,4})\b",
    ]
    normalized = normalize_ocr_text(text)
    for pattern in patterns:
        match = re.search(pattern, normalized)
        if match:
            return re.sub(r"\s+", " ", match.group(1)).strip()
    return ""


def guess_amounts_from_text(text: str) -> dict:
    normalized_lines = [normalize_ocr_text(line) for line in text.splitlines() if line.strip()]
    numeric_candidates = []

    for line in normalized_lines:
        for match in re.findall(r"\d+[.,]\d{2,3}", line):
            numeric_candidates.append((safe_float(match), line))

    litre = 0.0
    birim_fiyat = 0.0
    toplam = 0.0
    kdv_tutari = 0.0

    for value, line in numeric_candidates:
        if value <= 0:
            continue
        if ("LT" in line or "LITRE" in line) and not litre:
            litre = value
        elif any(token in line for token in ("FIYAT", "BIRIM")) and not birim_fiyat:
            birim_fiyat = value
        elif any(token in line for token in ("KDV", "VERGI")) and not kdv_tutari:
            kdv_tutari = value
        elif any(token in line for token in ("TOPLAM", "YOPLAM", "ODENECEK", "TUTAR", "NAKIT", "KAKIT", "WAKIT")):
            toplam = max(toplam, value)

    if not toplam:
        bigger_values = [value for value, _ in numeric_candidates if value > 20]
        if bigger_values:
            toplam = max(bigger_values)

    if not litre:
        litre_values = [value for value, line in numeric_candidates if 1 <= value <= 120 and ("LT" in line or value < toplam)]
        if litre_values:
            litre = litre_values[0]

    if not birim_fiyat and litre and toplam:
        birim_fiyat = round(toplam / litre, 3) if litre else 0.0

    return {
        "litre": litre,
        "birim_fiyat": birim_fiyat,
        "kdv_tutari": kdv_tutari,
        "toplam": toplam,
    }


def extract_payment_type(text: str) -> str:
    normalized = normalize_ocr_text(text)
    if "NAKIT" in normalized:
        return "Nakit"
    if any(token in normalized for token in ("KREDI KARTI", "POS", "BANKA KARTI", "VISA", "MASTERCARD", "K.KART", "KKART", "KART")):
        return "Kart"
    if "VERESIYE" in normalized:
        return "Veresiye"
    return ""


def extract_structured_fuel_fields(text: str) -> dict:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    normalized_lines = [normalize_ocr_text(line) for line in lines]

    found = {
        "vergi_no": "",
        "tarih": "",
        "saat": "",
        "fis_no": "",
        "plaka": "",
        "litre": "",
        "birim_fiyat": "",
        "kdv_tutari": "",
        "toplam": "",
        "odeme_tipi": "",
    }

    for line, normalized in zip(lines, normalized_lines):
        if not found["vergi_no"]:
            match = re.search(r"\b(?:VD|VKN|VERGI\s*NO)[: ]+(\d{10,11})\b", normalized)
            if match:
                found["vergi_no"] = match.group(1)

        if not found["tarih"]:
            match = re.search(r"TARIH[: ]+(\d{2}[./-]\d{2}[./-]\d{2,4})", normalized)
            if match:
                found["tarih"] = normalize_date(match.group(1))
            else:
                match = re.search(r"\b(\d{2}[./-]\d{2}[./-]\d{2,4})\b", normalized)
                if match:
                    found["tarih"] = normalize_date(match.group(1))

        if not found["saat"]:
            match = re.search(r"SAAT[: ]+(\d{2}[:.]\d{2})", normalized)
            if match:
                found["saat"] = normalize_time(match.group(1))
            else:
                match = re.search(r"\b(\d{2}[:.]\d{2})\b", normalized)
                if match:
                    found["saat"] = normalize_time(match.group(1))

        if not found["fis_no"]:
            match = re.search(r"FIS\s*NO[: ]*([A-Z0-9]+)", normalized)
            if match:
                found["fis_no"] = match.group(1)

        if not found["plaka"]:
            plate = extract_plate(normalized)
            if not plate:
                compact_match = re.search(r"(0[1-9]|[1-7][0-9]|8[01])[A-Z]{1,3}\d{2,4}", re.sub(r"[^A-Z0-9]", "", normalized))
                if compact_match:
                    plate = compact_match.group(0)
            if plate:
                found["plaka"] = plate

        if " LT " in f" {normalized} " and " X " in f" {normalized} ":
            match = re.search(r"(\d[\d.,]*)\s*LT\s*X\s*(\d[\d.,]*)", normalized)
            if match:
                found["litre"] = normalize_number_token(match.group(1))
                found["birim_fiyat"] = normalize_number_token(match.group(2))

        if not found["kdv_tutari"] and "KDV" in normalized:
            found["kdv_tutari"] = extract_line_amount(line)

        if not found["toplam"] and "TOPLAM" in normalized:
            found["toplam"] = extract_line_amount(line)

        if not found["odeme_tipi"] and any(token in normalized for token in ("K.KART", "KKART", "KART", "NAKIT")):
            found["odeme_tipi"] = extract_payment_type(normalized)

    return found


def extract_fuel_receipt_fields(text: str) -> dict:
    normalized_text = normalize_ocr_text(text)
    compact_text = compact_ocr_text(text)
    structured = extract_structured_fuel_fields(text)
    date_value = extract_first_match(
        [
            r"\b(\d{2}[./-]\d{2}[./-]\d{2,4})\b",
        ],
        normalized_text,
    )
    time_value = extract_first_match([r"\b(\d{2}[:.]\d{2})\b"], text)
    fis_no = extract_first_match(
        [
            r"F[Iİ]S(?:\s*NO)?[: ]+([A-Z0-9\-\/]+)",
            r"FIS\s*WQ\.?[: ]*([A-Z0-9\-\/]+)",
            r"FIS\s*NO\.?[: ]*([A-Z0-9\-\/]+)",
            r"BELGE\s*NO[: ]+([A-Z0-9\-\/]+)",
            r"NO[: ]+([A-Z0-9\-\/]{3,})",
            r"UTTS.*?([A-Z0-9]{3,})",
        ],
        normalized_text,
    )
    vergi_no = extract_first_match(
        [
            r"VN[: ]+(\d{10,11})",
            r"V(?:ERGI)?\.?\s*NO[: ]+(\d{10,11})",
            r"VKN[: ]+(\d{10,11})",
        ],
        normalized_text,
    )
    litre_match = extract_first_match(
        [
            r"L[İI]TRE[: ]+([\d.,]+)",
            r"([\d.,]+)\s*LT\s*X",
            r"([\d.,]+)\s*LT\b",
        ],
        normalized_text,
    )
    birim_fiyat_match = extract_first_match(
        [
            r"B[Iİ]R[Iİ]M\s*FIYAT[: ]+([\d.,]+)",
            r"L[İI]TRE\s*FIYAT[İI]?[: ]+([\d.,]+)",
            r"LT\s*X\s*([\d.,]+)",
        ],
        normalized_text,
    )
    kdv_match = extract_first_match(
        [
            r"KDV(?:\s*TUTAR[Iİ])?[: ]+([\d.,]+)",
            r"KD[VY][:. ]+([\d.,]+)",
        ],
        normalized_text,
    )
    toplam_match = extract_first_match(
        [
            r"TOPLAM(?:\s*TUTAR)?[: ]+([\d.,]+)",
            r"GENEL\s*TOPLAM[: ]+([\d.,]+)",
            r"ODENECEK[: ]+([\d.,]+)",
            r"TOPLAMTUTAR[: ]*([\d.,]+)",
            r"YOPLAM[: ]+([\d.,]+)",
            r"[NWK]AKIT[: ]+([\d.,]+)",
        ],
        normalized_text,
    )
    guessed_amounts = guess_amounts_from_text(text)
    plate = extract_plate(text)
    if not plate:
        compact_plate = re.search(r"(0[1-9]|[1-7][0-9]|8[01])[A-Z]{1,3}\d{2,4}", compact_text)
        if compact_plate:
            plate = compact_plate.group(0)

    return {
        "istasyon_adi": pick_station_name(text),
        "vergi_no": structured["vergi_no"] or vergi_no,
        "tarih": structured["tarih"] or normalize_date(date_value),
        "saat": structured["saat"] or normalize_time(time_value),
        "fis_no": structured["fis_no"] or fis_no,
        "plaka": structured["plaka"] or plate,
        "litre": structured["litre"] or litre_match or (f"{guessed_amounts['litre']:.3f}" if guessed_amounts["litre"] else ""),
        "birim_fiyat": structured["birim_fiyat"] or birim_fiyat_match or (f"{guessed_amounts['birim_fiyat']:.3f}" if guessed_amounts["birim_fiyat"] else ""),
        "kdv_tutari": structured["kdv_tutari"] or kdv_match or (f"{guessed_amounts['kdv_tutari']:.2f}" if guessed_amounts["kdv_tutari"] else ""),
        "toplam": structured["toplam"] or toplam_match or (f"{guessed_amounts['toplam']:.2f}" if guessed_amounts["toplam"] else ""),
        "odeme_tipi": structured["odeme_tipi"] or extract_payment_type(text),
        "aciklama": "Yakit fisi",
        "ocr_metin": text,
    }


def extract_fuel_receipt(filepath: Path) -> dict:
    text = extract_text_with_ocr(filepath)
    if not text.strip():
        raise HTTPException(status_code=422, detail="Yakit fisinden okunabilir metin cikarilamadi.")
    return extract_fuel_receipt_fields(text)


def build_claude_content(filepath: Path, prompt: str) -> list[dict]:
    mime_type = get_mime_type(filepath.name)
    b64 = file_to_base64(filepath)

    if mime_type == "application/pdf":
        file_block = {
            "type": "document",
            "source": {
                "type": "base64",
                "media_type": mime_type,
                "data": b64,
            },
        }
    else:
        file_block = {
            "type": "image",
            "source": {
                "type": "base64",
                "media_type": mime_type,
                "data": b64,
            },
        }

    return [
        file_block,
        {
            "type": "text",
            "text": prompt,
        },
    ]


def extract_with_claude(filepath: Path) -> dict:
    if not CLAUDE_API_KEY:
        raise HTTPException(status_code=500, detail="CLAUDE_API_KEY tanımlı değil.")
    model_name = normalize_model_name(CLAUDE_MODEL)

    prompt = """
Sen Türkiye'de çalışan deneyimli bir mali müşavir yardımcısısın.
Kullanıcının yüklediği belge bir perakende satış fişi fotoğrafıdır.
Fişteki alanları analiz et ve sadece aşağıdaki JSON yapısında cevap ver:
{
  "belge_tipi": "",
  "firma": "",
  "vergi_no": "",
  "tarih": "",
  "belge_no": "",
  "matrah": "",
  "kdv_orani": "",
  "kdv_tutari": "",
  "toplam": "",
  "aciklama": ""
}
Kurallar:
- Çıktıyı sadece JSON olarak ver
- Tarih formatı YYYY-MM-DD olsun
- Emin olmadığın alanları boş bırak
- Türkiye perakende satış fişlerine göre yorumla
- Belge tipi çoğunlukla "fiş" olsun
- Sayısal alanlarda sadece sayı döndür
"""
    payload = {
        "model": model_name,
        "max_tokens": 1024,
        "messages": [
            {
                "role": "user",
                "content": build_claude_content(filepath, prompt),
            }
        ],
    }
    headers = {
        "x-api-key": CLAUDE_API_KEY,
        "anthropic-version": "2023-06-01",
        "content-type": "application/json",
    }
    response = requests.post(
        "https://api.anthropic.com/v1/messages",
        headers=headers,
        json=payload,
        timeout=120,
    )
    try:
        response.raise_for_status()
    except requests.HTTPError as exc:
        detail = response.text[:500] if response.text else "AI isteği başarısız oldu."
        raise HTTPException(status_code=502, detail=detail) from exc

    data = response.json()
    text_output = ""
    for content in data.get("content", []):
        if content.get("type") == "text":
            text_output += content.get("text", "")
    if not text_output.strip():
        raise HTTPException(status_code=500, detail="AI cevabı alınamadı.")
    try:
        parsed = json.loads(text_output)
    except json.JSONDecodeError:
        cleaned = text_output.strip().removeprefix("```json").removeprefix("```").removesuffix("```").strip()
        parsed = json.loads(cleaned)
    return {
        "belge_tipi": parsed.get("belge_tipi", ""),
        "firma": parsed.get("firma", ""),
        "vergi_no": parsed.get("vergi_no", ""),
        "tarih": parsed.get("tarih", ""),
        "belge_no": parsed.get("belge_no", ""),
        "matrah": parsed.get("matrah", ""),
        "kdv_orani": parsed.get("kdv_orani", ""),
        "kdv_tutari": parsed.get("kdv_tutari", ""),
        "toplam": parsed.get("toplam", ""),
        "aciklama": parsed.get("aciklama", ""),
    }


@app.get("/health")
def health_check():
    return {"status": "ok"}


@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    template = jinja_env.get_template("index.html")
    return HTMLResponse(template.render(request=request))


@app.get("/fuel", response_class=HTMLResponse)
def fuel_home(request: Request):
    template = jinja_env.get_template("fuel.html")
    return HTMLResponse(template.render(request=request))


@app.post("/upload", response_class=HTMLResponse)
async def upload_document(
    request: Request,
    musteri: str = Form(...),
    file: list[UploadFile] = File(...),
):
    files = [item for item in file if item.filename]
    if not files:
        raise HTTPException(status_code=400, detail="Dosya seçilmedi.")

    saved_count = 0

    for index, upload in enumerate(files, start=1):
        mime_type = get_mime_type(upload.filename)
        if not is_image_mime_type(mime_type):
            raise HTTPException(status_code=400, detail="Sadece perakende satış fişi fotoğrafları yükleyebilirsiniz.")

        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        safe_name = f"{timestamp}_{index}_{sanitize_upload_filename(upload.filename)}"
        file_path = UPLOAD_DIR / safe_name

        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(upload.file, buffer)

        try:
            public_file_url = upload_to_supabase_storage(file_path, safe_name, mime_type)
            ai_result = extract_with_claude(file_path)
            db = SessionLocal()
            try:
                record = Record(
                    musteri=musteri,
                    belge_tipi=ai_result.get("belge_tipi", "") or "fiş",
                    firma=ai_result.get("firma", ""),
                    vergi_no=ai_result.get("vergi_no", ""),
                    tarih=ai_result.get("tarih", ""),
                    belge_no=ai_result.get("belge_no", ""),
                    matrah=safe_float(ai_result.get("matrah", "")),
                    kdv_orani=ai_result.get("kdv_orani", ""),
                    kdv_tutari=safe_float(ai_result.get("kdv_tutari", "")),
                    toplam=safe_float(ai_result.get("toplam", "")),
                    aciklama=ai_result.get("aciklama", ""),
                    dosya_adi=safe_name,
                    dosya_yolu=public_file_url,
                    durum="Onaylandı",
                )
                db.add(record)
                db.commit()
                saved_count += 1
            finally:
                db.close()
        except Exception:
            if file_path.exists():
                file_path.unlink()
            raise
        finally:
            if file_path.exists() and not should_keep_local_upload():
                file_path.unlink()

    return RedirectResponse(url=f"/records?uploaded={saved_count}", status_code=303)


@app.post("/fuel/upload", response_class=HTMLResponse)
async def upload_fuel_receipt(
    request: Request,
    musteri: str = Form(...),
    file: UploadFile = File(...),
):
    if not file.filename:
        raise HTTPException(status_code=400, detail="Dosya secilmedi.")

    mime_type = get_mime_type(file.filename)
    if not is_image_mime_type(mime_type):
        raise HTTPException(status_code=400, detail="Yakit fisi modu icin sadece resim yukleyebilirsiniz.")

    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    safe_name = f"fuel_{timestamp}_{sanitize_upload_filename(file.filename)}"
    file_path = UPLOAD_DIR / safe_name

    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    try:
        public_file_url = upload_to_supabase_storage(file_path, safe_name, mime_type)
        fuel_result = extract_fuel_receipt(file_path)
    except Exception:
        if file_path.exists():
            file_path.unlink()
        raise
    finally:
        if file_path.exists() and not should_keep_local_upload():
            file_path.unlink()

    template = jinja_env.get_template("fuel_review.html")
    return HTMLResponse(
        template.render(
            request=request,
            musteri=musteri,
            dosya_adi=safe_name,
            dosya_yolu=public_file_url,
            dosya_url=public_file_url,
            result=fuel_result,
        )
    )


@app.post("/save")
async def save_record(
    musteri: str = Form(...),
    belge_tipi: str = Form(""),
    firma: str = Form(""),
    vergi_no: str = Form(""),
    tarih: str = Form(""),
    belge_no: str = Form(""),
    matrah: str = Form(""),
    kdv_orani: str = Form(""),
    kdv_tutari: str = Form(""),
    toplam: str = Form(""),
    aciklama: str = Form(""),
    dosya_adi: str = Form(""),
    dosya_yolu: str = Form(""),
    durum: str = Form("Onaylandı"),
):
    db = SessionLocal()
    try:
        record = Record(
            musteri=musteri,
            belge_tipi=belge_tipi,
            firma=firma,
            vergi_no=vergi_no,
            tarih=tarih,
            belge_no=belge_no,
            matrah=safe_float(matrah),
            kdv_orani=kdv_orani,
            kdv_tutari=safe_float(kdv_tutari),
            toplam=safe_float(toplam),
            aciklama=aciklama,
            dosya_adi=dosya_adi,
            dosya_yolu=dosya_yolu,
            durum=durum,
        )
        db.add(record)
        db.commit()
    finally:
        db.close()
    return RedirectResponse(url="/records", status_code=303)


@app.post("/fuel/save")
async def save_fuel_record(
    musteri: str = Form(...),
    istasyon_adi: str = Form(""),
    vergi_no: str = Form(""),
    tarih: str = Form(""),
    saat: str = Form(""),
    fis_no: str = Form(""),
    plaka: str = Form(""),
    litre: str = Form(""),
    birim_fiyat: str = Form(""),
    kdv_tutari: str = Form(""),
    toplam: str = Form(""),
    odeme_tipi: str = Form(""),
    aciklama: str = Form(""),
    ocr_metin: str = Form(""),
    dosya_adi: str = Form(""),
    dosya_yolu: str = Form(""),
    durum: str = Form("Onaylandi"),
):
    db = SessionLocal()
    try:
        record = FuelRecord(
            musteri=musteri,
            istasyon_adi=istasyon_adi,
            vergi_no=vergi_no,
            tarih=tarih,
            saat=saat,
            fis_no=fis_no,
            plaka=plaka,
            litre=safe_float(litre),
            birim_fiyat=safe_float(birim_fiyat),
            kdv_tutari=safe_float(kdv_tutari),
            toplam=safe_float(toplam),
            odeme_tipi=odeme_tipi,
            aciklama=aciklama,
            ocr_metin=ocr_metin,
            dosya_adi=dosya_adi,
            dosya_yolu=dosya_yolu,
            durum=durum,
        )
        db.add(record)
        db.commit()
    finally:
        db.close()
    return RedirectResponse(url="/fuel-records", status_code=303)


@app.get("/records", response_class=HTMLResponse)
def list_records(request: Request):
    db = SessionLocal()
    try:
        records = db.query(Record).order_by(Record.created_at.desc()).all()
    finally:
        db.close()
    for record in records:
        record.preview_url = build_public_file_url(record.dosya_yolu or record.dosya_adi or "")
    template = jinja_env.get_template("records.html")
    return HTMLResponse(template.render(request=request, records=records, uploaded=request.query_params.get("uploaded", "")))


@app.get("/fuel-records", response_class=HTMLResponse)
def list_fuel_records(request: Request):
    db = SessionLocal()
    try:
        records = db.query(FuelRecord).order_by(FuelRecord.created_at.desc()).all()
    finally:
        db.close()
    for record in records:
        record.preview_url = build_public_file_url(record.dosya_yolu or record.dosya_adi or "")
    template = jinja_env.get_template("fuel_records.html")
    return HTMLResponse(template.render(request=request, records=records))


@app.get("/export")
def export_excel():
    db = SessionLocal()
    try:
        records = db.query(Record).order_by(Record.created_at.asc()).all()
    finally:
        db.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "Faturalar"
    headers = [
        "SIRANO",
        "Tarih",
        "EVRAKNO",
        "BelgeNo",
        "EVRAKTURU",
        "CARIKODU",
        "CARIADI",
        "STOKKODU",
        "STOKADI",
        "Birim",
        "DOVIZC",
        "Miktar",
        "BIRIMFIYATTL",
        "BRFD",
        "KDVY",
        "TUTARTL",
        "TUTARD",
        "ACIKLAMA",
        "TEVKORANI",
        "STOPAJORANI",
        "OTVY",
        "OZELKOD1",
        "OZELKOD2",
        "KASAKODU",
        "BANKAKODU",
        "INDY",
        "INDTL",
        "FATURACK",
        "IRSALIYENO",
        "IRSALIYETAR",
        "FATURAOK1",
        "FATURAOK2",
        "CARIADRESI",
        "Semt",
        "SEHIR",
        "VERGINO",
        "VERGIDAIRESI",
        "CARIMUHKODU",
        "HAREKETTIPI",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAD3")

    evrak_counters = {}
    for index, r in enumerate(records, start=1):
        receipt_date = parse_record_date(r.tarih)
        evrak_no = build_sales_receipt_no(receipt_date)
        evrak_counters[evrak_no] = evrak_counters.get(evrak_no, 0) + 1
        net_amount = kdv_excluded_amount(r)
        kdv_rate = safe_kdv_rate(r.kdv_orani, safe_float(r.matrah), safe_float(r.kdv_tutari))
        ws.append(
            [
                index,
                receipt_date,
                evrak_no,
                f"{evrak_no}-{evrak_counters[evrak_no]}",
                "Alış Faturası",
                "329.04.999",
                "",
                "770,02,004",
                "",
                "Adet",
                "",
                1,
                net_amount,
                "",
                kdv_rate,
                net_amount,
                "",
                r.firma or r.aciklama or "",
                "---",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "Açık",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "Masraf",
            ]
        )
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[1].number_format = "dd.mm.yyyy"
        row[12].number_format = "#,##0.00"
        row[14].number_format = "0.##"
        row[15].number_format = "#,##0.00"

    widths = {
        "A": 10,
        "B": 12,
        "C": 13,
        "D": 16,
        "E": 16,
        "F": 14,
        "H": 14,
        "J": 10,
        "L": 10,
        "M": 14,
        "O": 10,
        "P": 14,
        "R": 28,
        "AB": 12,
        "AM": 14,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    export_path = EXPORT_DIR / f"satis_fisleri_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(export_path)
    return FileResponse(
        path=export_path,
        filename=export_path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/fuel-export")
def export_fuel_excel():
    db = SessionLocal()
    try:
        records = db.query(FuelRecord).order_by(FuelRecord.created_at.desc()).all()
    finally:
        db.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Yakit Fisleri"
    headers = [
        "Musteri",
        "Istasyon",
        "Vergi No",
        "Tarih",
        "Saat",
        "Fis No",
        "Plaka",
        "Litre",
        "Birim Fiyat",
        "KDV Tutari",
        "Toplam",
        "Odeme Tipi",
        "Aciklama",
        "Dosya Adi",
        "Durum",
        "Kayit Tarihi",
    ]
    ws.append(headers)
    for r in records:
        ws.append(
            [
                r.musteri,
                r.istasyon_adi,
                r.vergi_no,
                r.tarih,
                r.saat,
                r.fis_no,
                r.plaka,
                r.litre,
                r.birim_fiyat,
                r.kdv_tutari,
                r.toplam,
                r.odeme_tipi,
                r.aciklama,
                r.dosya_adi,
                r.durum,
                r.created_at.strftime("%Y-%m-%d %H:%M"),
            ]
        )
    export_path = EXPORT_DIR / f"yakit_fisleri_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(export_path)
    return FileResponse(
        path=export_path,
        filename=export_path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
